import io
import logging
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# Conditional formatting fills
RED_FILL = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
GREEN_FILL = PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid")

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _style_header(ws, col_count: int):
    """Apply header styling to first row."""
    for col in range(1, col_count + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = THIN_BORDER


def _add_score_formatting(ws, col_letter: str, min_row: int, max_row: int):
    """Add conditional formatting for score columns (red <4, yellow 4-6, green >6)."""
    cell_range = f"{col_letter}{min_row}:{col_letter}{max_row}"
    ws.conditional_formatting.add(
        cell_range,
        CellIsRule(operator="lessThan", formula=["4"], fill=RED_FILL),
    )
    ws.conditional_formatting.add(
        cell_range,
        CellIsRule(operator="between", formula=["4", "6"], fill=YELLOW_FILL),
    )
    ws.conditional_formatting.add(
        cell_range,
        CellIsRule(operator="greaterThan", formula=["6"], fill=GREEN_FILL),
    )


def _auto_width(ws, col_count: int, max_width: int = 50):
    """Auto-adjust column widths."""
    for col in range(1, col_count + 1):
        max_len = 0
        letter = get_column_letter(col)
        for row in ws.iter_rows(min_col=col, max_col=col, values_only=False):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max_len + 2, max_width)


def generate_report(
    summary: dict[str, Any],
    employee_rankings: list[dict[str, Any]],
    calls: list[dict[str, Any]],
    transcripts: list[dict[str, Any]],
    deal_analyses: list[dict[str, Any]],
    detailed_analyses: list[dict[str, Any]],
    recommendations: list[dict[str, Any]],
) -> bytes:
    """Generate the 7-sheet Excel report."""
    wb = Workbook()

    # ========== Sheet 1: Сводка ==========
    ws1 = wb.active
    ws1.title = "Сводка"
    headers = ["Показатель", "Значение"]
    ws1.append(headers)
    _style_header(ws1, len(headers))

    rows = [
        ("Период анализа", summary.get("period", "")),
        ("Количество звонков", summary.get("total_calls", 0)),
        ("Количество сотрудников", summary.get("total_employees", 0)),
        ("Количество сделок", summary.get("total_deals", 0)),
        ("Средняя оценка по отделу", summary.get("avg_score", "-")),
        ("Средняя длительность звонка (сек)", summary.get("avg_duration", "-")),
        ("Звонки с фиксацией следующего шага", summary.get("calls_with_next_step", 0)),
        ("Звонки без следующего шага", summary.get("calls_without_next_step", 0)),
    ]
    for row in rows:
        ws1.append(row)

    # Top problems
    ws1.append(("", ""))
    ws1.append(("Топ-3 проблемы", ""))
    for i, problem in enumerate(summary.get("top_problems", [])[:3], 1):
        ws1.append((f"  {i}.", problem))

    ws1.append(("", ""))
    ws1.append(("Топ-3 рекомендации", ""))
    for i, rec in enumerate(summary.get("top_recommendations", [])[:3], 1):
        ws1.append((f"  {i}.", rec))

    _auto_width(ws1, 2)

    # ========== Sheet 2: Рейтинг менеджеров ==========
    ws2 = wb.create_sheet("Рейтинг менеджеров")
    headers = [
        "ФИО", "Звонков", "Средняя оценка", "Приветствие",
        "Потребности", "Презентация", "Возражения", "Закрытие",
        "Инициатива", "Приоритет обучения",
    ]
    ws2.append(headers)
    _style_header(ws2, len(headers))

    for emp in employee_rankings:
        ws2.append([
            emp.get("name", ""),
            emp.get("total_calls", 0),
            emp.get("avg_overall", ""),
            emp.get("avg_greeting", ""),
            emp.get("avg_needs_discovery", ""),
            emp.get("avg_presentation", ""),
            emp.get("avg_objection_handling", ""),
            emp.get("avg_closing", ""),
            emp.get("avg_initiative", ""),
            emp.get("priority_training", ""),
        ])

    max_row = ws2.max_row
    for col_letter in ["C", "D", "E", "F", "G", "H", "I"]:
        _add_score_formatting(ws2, col_letter, 2, max_row)
    _auto_width(ws2, len(headers))

    # ========== Sheet 3: Все звонки ==========
    ws3 = wb.create_sheet("Все звонки")
    headers = [
        "Дата", "Менеджер", "Направление", "Клиент (телефон)",
        "Длительность (сек)", "Сделка", "Общая оценка",
        "Кто ведёт", "Следующий шаг",
    ]
    ws3.append(headers)
    _style_header(ws3, len(headers))

    for call in calls:
        ws3.append([
            str(call.get("call_date", "")),
            call.get("employee_name", ""),
            call.get("direction", ""),
            call.get("phone_number", ""),
            call.get("duration_sec", 0),
            call.get("deal_name", ""),
            call.get("overall_score", ""),
            call.get("who_leads", ""),
            "Да" if call.get("next_step_agreed") else "Нет",
        ])

    max_row = ws3.max_row
    _add_score_formatting(ws3, "G", 2, max_row)
    _auto_width(ws3, len(headers))
    ws3.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max_row}"

    # ========== Sheet 4: Транскрипции ==========
    ws4 = wb.create_sheet("Транскрипции")
    headers = ["ID звонка", "Менеджер", "Сделка", "Транскрипт"]
    ws4.append(headers)
    _style_header(ws4, len(headers))

    for t in transcripts:
        ws4.append([
            t.get("call_id", ""),
            t.get("employee_name", ""),
            t.get("deal_name", ""),
            t.get("text", ""),
        ])

    _auto_width(ws4, len(headers))

    # ========== Sheet 5: Анализ по сделкам ==========
    ws5 = wb.create_sheet("Анализ по сделкам")
    headers = [
        "Сделка", "Менеджер", "Кол-во звонков", "Суммарная длительность (сек)",
        "Стадия", "Средняя оценка", "Проблемы",
    ]
    ws5.append(headers)
    _style_header(ws5, len(headers))

    for deal in deal_analyses:
        ws5.append([
            deal.get("deal_name", ""),
            deal.get("employee_name", ""),
            deal.get("call_count", 0),
            deal.get("total_duration", 0),
            deal.get("stage", ""),
            deal.get("avg_score", ""),
            "; ".join(deal.get("problems", [])),
        ])

    _auto_width(ws5, len(headers))

    # ========== Sheet 6: Детальный анализ ==========
    ws6 = wb.create_sheet("Детальный анализ")
    headers = [
        "ID звонка", "Менеджер", "Приветствие", "Потребности",
        "Презентация", "Возражения", "Закрытие", "Инициатива",
        "Общая", "Сильные стороны", "Слабые стороны",
        "Рекомендации", "Пропущенные шаги скрипта", "CRM-заметка",
    ]
    ws6.append(headers)
    _style_header(ws6, len(headers))

    for a in detailed_analyses:
        ws6.append([
            a.get("call_id", ""),
            a.get("employee_name", ""),
            a.get("greeting_score", ""),
            a.get("needs_discovery", ""),
            a.get("presentation_score", ""),
            a.get("objection_handling", ""),
            a.get("closing_score", ""),
            a.get("initiative_score", ""),
            a.get("overall_score", ""),
            "\n".join(a.get("strengths", [])),
            "\n".join(a.get("weaknesses", [])),
            "\n".join(a.get("recommendations", [])),
            "\n".join(a.get("missed_script_steps", [])),
            a.get("crm_note_suggestion", ""),
        ])

    max_row = ws6.max_row
    for col_letter in ["C", "D", "E", "F", "G", "H", "I"]:
        _add_score_formatting(ws6, col_letter, 2, max_row)
    _auto_width(ws6, len(headers))

    # ========== Sheet 7: Рекомендации ==========
    ws7 = wb.create_sheet("Рекомендации")
    headers = [
        "Приоритет", "Категория", "Проблема",
        "Затронутые менеджеры", "Рекомендация", "Ожидаемый эффект",
    ]
    ws7.append(headers)
    _style_header(ws7, len(headers))

    for rec in recommendations:
        ws7.append([
            rec.get("priority", ""),
            rec.get("category", ""),
            rec.get("problem", ""),
            rec.get("affected_employees", ""),
            rec.get("recommendation", ""),
            rec.get("expected_effect", ""),
        ])

    _auto_width(ws7, len(headers))

    # Save to bytes
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()
